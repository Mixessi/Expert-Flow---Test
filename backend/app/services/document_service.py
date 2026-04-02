import os
import uuid

from sqlalchemy.ext.asyncio import AsyncSession

from app.config import settings
from app.models.reference_document import ReferenceDocument


async def save_document(
    db: AsyncSession,
    interview_id: uuid.UUID,
    filename: str,
    file_content: bytes,
    file_type: str,
) -> ReferenceDocument:
    """Save uploaded document and extract text content."""
    # Ensure upload directory exists
    upload_dir = os.path.join(settings.upload_dir, str(interview_id))
    os.makedirs(upload_dir, exist_ok=True)

    # Save file
    file_path = os.path.join(upload_dir, filename)
    with open(file_path, "wb") as f:
        f.write(file_content)

    # Extract text
    extracted_text = extract_text(file_content, file_type, filename)

    doc = ReferenceDocument(
        interview_id=interview_id,
        filename=filename,
        file_type=file_type,
        file_size=len(file_content),
        extracted_text=extracted_text,
        file_path=file_path,
    )
    db.add(doc)
    await db.commit()
    await db.refresh(doc)
    return doc


def extract_text(content: bytes, file_type: str, filename: str) -> str:
    """Extract text from document based on file type."""
    lower_name = filename.lower()

    if lower_name.endswith(".pdf"):
        return _extract_pdf(content)
    elif lower_name.endswith((".docx", ".doc")):
        return _extract_docx(content)
    elif lower_name.endswith((".txt", ".md", ".csv")):
        return content.decode("utf-8", errors="replace")
    else:
        return content.decode("utf-8", errors="replace")


def _extract_pdf(content: bytes) -> str:
    try:
        import io
        from PyPDF2 import PdfReader
        reader = PdfReader(io.BytesIO(content))
        texts = []
        for page in reader.pages:
            text = page.extract_text()
            if text:
                texts.append(text)
        return "\n".join(texts)
    except Exception as e:
        return f"[PDF extraction failed: {e}]"


def _extract_docx(content: bytes) -> str:
    try:
        import io
        from docx import Document
        doc = Document(io.BytesIO(content))
        return "\n".join(para.text for para in doc.paragraphs if para.text.strip())
    except Exception as e:
        return f"[DOCX extraction failed: {e}]"


async def get_documents_text(db: AsyncSession, interview_id: uuid.UUID) -> str:
    """Get concatenated text of all reference documents for an interview."""
    from sqlalchemy import select
    result = await db.execute(
        select(ReferenceDocument).where(ReferenceDocument.interview_id == interview_id)
    )
    docs = result.scalars().all()
    texts = []
    for doc in docs:
        if doc.extracted_text:
            texts.append(f"--- 文档: {doc.filename} ---\n{doc.extracted_text}")
    return "\n\n".join(texts)
